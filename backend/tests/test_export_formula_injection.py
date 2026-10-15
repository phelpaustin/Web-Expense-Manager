"""Release 13: exported CSV/XLSX cannot carry spreadsheet formulas."""
import csv
import io

import pytest
from openpyxl import load_workbook

from app.core.parsing import escape_formula, unescape_formula

EVIL = ['=HYPERLINK("http://evil.example","x")', "+1+1", "-2+3", "@SUM(1+1)", "\t=1+1", "=cmd|' /C calc'!A0"]


@pytest.mark.parametrize("value", EVIL)
def test_escape_prefixes_quote_and_roundtrips(value):
    escaped = escape_formula(value)
    assert escaped == "'" + value
    assert unescape_formula(escaped) == value


@pytest.mark.parametrize("value", ["Lidl", "Coffee - large", "50% off", "", "'quoted'", "ICA @ Maxi", 12.5, None])
def test_safe_values_untouched(value):
    assert escape_formula(value) == value
    assert unescape_formula(value) == value


def _add(client, h, **kw):
    body = {"date": "2026-10-01", "category": "Food", "description": "x", "shop": "s", "amount": 100, "currency": "SEK"}
    body.update(kw)
    r = client.post("/api/expenses", json=body, headers=h)
    assert r.status_code in (200, 201), r.text


def test_csv_export_neutralises_text_cells(client, make_user):
    _, h = make_user()
    _add(client, h, shop=EVIL[0], description=EVIL[1], category="=Cat", subcategory="@sub")
    _add(client, h, shop="Lidl", description="Milk", amount=25)
    r = client.get("/api/expenses/export?format=csv", headers=h)
    rows = list(csv.DictReader(io.StringIO(r.text)))
    evil = next(x for x in rows if x["Item"].endswith("+1+1"))
    assert evil["Shop"] == "'" + EVIL[0] and evil["Item"] == "'+1+1"
    assert evil["Category"] == "'=Cat" and evil["Subcategory"] == "'@sub"
    for row in rows:
        for col in ("Category", "Subcategory", "Item", "Shop", "Brand", "QuantityUnit", "Currency"):
            assert not row[col].startswith(("=", "+", "-", "@", "\t")), (col, row[col])
    plain = next(x for x in rows if x["Shop"] == "Lidl")
    assert plain["PricePaid"] == "25.00" or float(plain["PricePaid"]) == 25  # numbers stay numbers


def test_xlsx_export_neutralises_text_cells(client, make_user):
    _, h = make_user()
    _add(client, h, shop=EVIL[0], description="-2+3")
    r = client.get("/api/expenses/export?format=xlsx", headers=h)
    ws = load_workbook(io.BytesIO(r.content)).active
    header = [c.value for c in ws[1]]
    row = [c.value for c in ws[2]]
    cells = dict(zip(header, row))
    assert cells["Shop"] == "'" + EVIL[0] and cells["Item"] == "'-2+3"
    for cell in ws[2]:
        assert cell.data_type != "f", f"formula cell {cell.coordinate}"


def test_reimporting_our_export_restores_original_text(client, make_user):
    _, h = make_user()
    _add(client, h, shop=EVIL[0], description="-2+3", category="Food")
    exported = client.get("/api/expenses/export?format=csv", headers=h).content
    _, h2 = make_user()
    r = client.post("/api/expenses/import", files={"file": ("expenses.csv", exported, "text/csv")}, headers=h2)
    assert r.status_code == 200 and r.json()["added"] == 1, r.text
    items = client.get("/api/expenses", headers=h2).json()
    items = items["items"] if isinstance(items, dict) and "items" in items else items
    assert items[0]["shop"] == EVIL[0] and items[0]["description"] == "-2+3"
