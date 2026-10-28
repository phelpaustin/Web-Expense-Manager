"""The Excel and PDF files built from a report."""
import datetime as dt
import io
import re
import shutil
import subprocess
from decimal import Decimal as D

import openpyxl
import pytest

from app.core import report_files as rf
from app.logic import report as rp

TODAY = dt.date(2026, 10, 28)


def e(date, amount, category="Food", description="x", shop=""):
    return {"id": 1, "date": dt.date.fromisoformat(date), "category": category, "description": description, "shop": shop, "amount": D(str(amount))}


ROWS = [
    e("2026-01-10", 100, "Food", "Milk", "ICA"), e("2026-02-11", 300, "Rent", "Rent", "Landlord"),
    e("2026-02-12", 50, "Food", "Bread", "ICA"), e("2026-10-05", 70, "Fun", "Cinema", "SF"),
]
INCOME = [{"date": dt.date(2026, 1, 25), "amount": D(1000)}, {"date": dt.date(2026, 2, 25), "amount": D(1000)}]


def year_report(rows=ROWS):
    return rp.build_report(rows, INCOME, 2026, None, today=TODAY)


def month_report(rows=ROWS):
    return rp.build_report(rows, INCOME, 2026, 2, today=TODAY)


def book(data):
    return openpyxl.load_workbook(io.BytesIO(data))


def sheet_rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def pdf_text(data):
    if not shutil.which("pdftotext"):
        pytest.skip("pdftotext is not installed")
    return subprocess.run(["pdftotext", "-layout", "-", "-"], input=data, capture_output=True, check=True).stdout.decode()


# ── names ──

def test_file_names():
    assert rf.file_stem(year_report()) == "report-2026" and rf.file_stem(month_report()) == "report-2026-02"


# ── Excel ──

def test_a_year_workbook_has_the_expected_sheets():
    wb = book(rf.report_xlsx(year_report(), "SEK"))
    assert wb.sheetnames == ["Summary", "Categories", "By month", "Category by month", "Top shops", "Largest expenses"]


def test_a_month_workbook_is_by_day_and_has_no_category_by_month_sheet():
    wb = book(rf.report_xlsx(month_report(), "SEK"))
    assert wb.sheetnames == ["Summary", "Categories", "By day", "Top shops", "Largest expenses"]
    assert len(sheet_rows(wb["By day"])) == 1 + 28


def test_the_summary_sheet_holds_the_totals():
    ws = book(rf.report_xlsx(year_report(), "SEK"))["Summary"]
    values = {r[0]: r[1] for r in sheet_rows(ws) if r[0]}
    assert values["Spending report 2026"] is None and values["All amounts in SEK"] is None
    assert values["Total spent"] == 520 and values["Number of expenses"] == 4
    assert values["Income"] == 2000 and values["Net (income minus spending)"] == 1480
    assert values["Savings rate"] == pytest.approx(0.74)
    assert values["Spent in 2025"] == 0 and values["Change in percent"] == "–"
    assert "Average per month" in values and values["Highest month: February"] == 350


def test_a_month_summary_has_no_monthly_average_and_names_a_day():
    values = {r[0]: r[1] for r in sheet_rows(book(rf.report_xlsx(month_report(), "SEK"))["Summary"]) if r[0]}
    assert "Average per month" not in values and values["Highest day: 11 February"] == 300
    assert values["Spent in January 2026"] == 100 and values["Change in percent"] == pytest.approx(2.5)


def test_amounts_are_real_numbers_with_a_money_format_and_shares_are_percent():
    ws = book(rf.report_xlsx(year_report(), "SEK"))["Categories"]
    rent = sheet_rows(ws)[1]
    assert rent[:4] == ["Rent", 300, pytest.approx(0.577), 1]            # the report rounds a share to one decimal
    assert ws["B2"].number_format == "#,##0.00" and ws["C2"].number_format == "0.0%"
    assert ws["B1"].value == "Amount (SEK)"
    assert sheet_rows(ws)[-1][:2] == ["Total", 520] and sheet_rows(ws)[-1][3] == 4


def test_the_category_and_period_sheets_each_have_a_native_chart():
    wb = book(rf.report_xlsx(year_report(), "SEK"))
    assert len(wb["Categories"]._charts) == 1 and len(wb["By month"]._charts) == 1
    assert len(wb["Summary"]._charts) == 0 and len(wb["Top shops"]._charts) == 0


def test_the_category_by_month_sheet_adds_up():
    ws = book(rf.report_xlsx(year_report(), "SEK"))["Category by month"]
    rows = sheet_rows(ws)
    assert rows[0][0] == "Category" and rows[0][1:4] == ["Jan", "Feb", "Mar"] and rows[0][-1] == "Total"
    rent = next(r for r in rows if r[0] == "Rent")
    assert rent[2] == 300 and rent[-1] == 300
    assert rows[-1][0] == "Total" and rows[-1][-1] == 520 and rows[-1][2] == 350


def test_dates_are_real_dates_and_text_cannot_run_as_a_formula():
    rows = [e("2026-02-01", 10, "=SUM(A1)", "=HYPERLINK(\"http://evil\")", "+cmd|calc"), e("2026-02-02", 5, "Food", "ok", "@x")]
    wb = book(rf.report_xlsx(rp.build_report(rows, [], 2026, None, today=TODAY), "SEK"))
    largest = sheet_rows(wb["Largest expenses"])
    assert isinstance(largest[1][0], dt.datetime) and largest[1][0].date() == dt.date(2026, 2, 1)
    evil = largest[1]
    assert evil[1].startswith("'=") and evil[2].startswith("'=") and evil[3].startswith("'+")
    assert sheet_rows(wb["Top shops"])[1][0].startswith("'+")
    assert any(r[0].startswith("'=") for r in sheet_rows(wb["Categories"])[1:-1])
    for ws in wb:                                       # and no cell anywhere is stored as a formula
        for row in ws.iter_rows():
            for c in row:
                assert c.data_type != "f", (ws.title, c.coordinate, c.value)


def test_an_empty_report_still_makes_a_valid_workbook():
    empty = rp.build_report([], [], 2026, None, today=TODAY)
    wb = book(rf.report_xlsx(empty, "SEK"))
    assert len(wb["Categories"]._charts) == 0 and sheet_rows(wb["Categories"]) == [["Category", "Amount (SEK)", "Share", "Expenses"]]
    assert sheet_rows(wb["Largest expenses"])[1:] == []
    book(rf.report_xlsx(rp.build_report([], [], 2026, 2, today=TODAY), "SEK"))


def test_many_categories_all_get_a_row():
    rows = [e("2026-02-01", 10 + i, f"Cat {i}") for i in range(40)]
    ws = book(rf.report_xlsx(rp.build_report(rows, [], 2026, None, today=TODAY), "SEK"))["Categories"]
    assert len(sheet_rows(ws)) == 1 + 40 + 1


# ── PDF ──

def test_a_pdf_is_a_pdf():
    for report in (year_report(), month_report()):
        data = rf.report_pdf(report, "SEK", TODAY)
        assert data.startswith(b"%PDF-") and data.rstrip().endswith(b"%%EOF") and len(data) > 2000


def test_the_pdf_says_what_the_report_says():
    text = pdf_text(rf.report_pdf(year_report(), "SEK", TODAY))
    for expected in ("Spending report 2026", "All amounts in SEK", "created 28 Oct 2026", "Total spent", "520.00", "Income", "2,000.00",
                     "74.0%", "Spending by category", "Rent", "300.00", "57.7%", "Spending by month", "Top shops", "Landlord", "Largest expenses", "SF"):
        assert expected in text, expected
    month = pdf_text(rf.report_pdf(month_report(), "SEK", TODAY))
    assert "Spending report February 2026" in month and "Spending by day" in month and "Highest day" in month and "+250.0%" in month


def test_a_pdf_without_income_leaves_out_the_income_lines():
    text = pdf_text(rf.report_pdf(rp.build_report(ROWS, [], 2026, None, today=TODAY), "SEK", TODAY))
    assert "Savings rate" not in text and "Net (income" not in text and "Total spent" in text


def test_an_empty_pdf_says_so():
    assert "No spending was recorded in this period." in pdf_text(rf.report_pdf(rp.build_report([], [], 2026, None, today=TODAY), "SEK", TODAY))


def test_characters_the_pdf_font_cannot_draw_do_not_break_it():
    rows = [e("2026-02-01", 10, "日本語", "Ωmega <b>&</b> åäö", "Café ☕")]
    data = rf.report_pdf(rp.build_report(rows, [], 2026, None, today=TODAY), "SEK", TODAY)
    text = pdf_text(data)
    assert "åäö" in text and "Café" in text and "<b>" in text and "???" in text


def test_pdf_text_helpers():
    assert rf._pdf_safe("Åäö é") == "Åäö é" and rf._pdf_safe("日本") == "??" and rf._pdf_safe(None) == ""
    assert rf._money(D("1234567.8")) == "1,234,567.80" and rf._money(D("-5")) == "-5.00"
    assert rf._short("x" * 40).endswith("…") and len(rf._short("x" * 40)) == rf.NAME_LIMIT and rf._short("short") == "short"


def test_a_big_report_runs_over_several_pages():
    rows = [e(f"2026-{1 + i % 12:02d}-{1 + i % 27:02d}", 10 + i, f"Category {i}", f"item {i}", f"Shop {i}") for i in range(80)]
    data = rf.report_pdf(rp.build_report(rows, [], 2026, None, today=TODAY), "SEK", TODAY)
    assert len(re.findall(rb"/Type\s*/Page\b", data)) >= 3
    text = pdf_text(data)
    assert "Other (65)" in text and "Category 79" in text           # the chart groups the tail, the table lists every category


def test_a_negative_total_does_not_break_the_charts():
    rows = [e("2026-02-01", -50, "Food"), e("2026-03-01", 20, "Rent")]
    assert rf.report_pdf(rp.build_report(rows, [], 2026, None, today=TODAY), "SEK", TODAY).startswith(b"%PDF")
    assert rf.report_pdf(rp.build_report([e("2026-02-01", -50, "Food")], [], 2026, 2, today=TODAY), "SEK", TODAY).startswith(b"%PDF")
