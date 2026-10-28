import datetime as dt
import io

import openpyxl

URL = "/api/reports"
THIS_YEAR = dt.date.today().year


def add(client, h, amount, when, category="Food", description="x", shop="", **extra):
    r = client.post("/api/expenses", json={"date": when, "category": category, "description": description, "amount": str(amount),
                                           "shop": shop, "currency": "SEK", **extra}, headers=h)
    assert r.status_code in (200, 201), r.text
    return r.json()


def get(client, h, path=URL, **params):
    r = client.get(path, params=params, headers=h)
    assert r.status_code == 200, r.text
    return r


def seed(client, h):
    add(client, h, 100, "2025-12-20", "Food", "Old", "ICA")
    add(client, h, 300, "2026-02-11", "Rent", "Rent", "Landlord")
    add(client, h, 50, "2026-02-12", "Food", "Bread", "ICA")
    add(client, h, 70, "2026-03-05", "Fun", "Cinema", "SF")
    r = client.post("/api/income", json={"date": "2026-02-25", "amount": "1000", "source": "Job"}, headers=h)
    assert r.status_code == 201, r.text


def test_a_year_report(client, make_user):
    _, h = make_user()
    seed(client, h)
    r = get(client, h, year=2026).json()
    assert (r["kind"], r["label"], r["year"], r["month"]) == ("year", "2026", 2026, None)
    assert r["totals"]["spent"] == 420 and r["totals"]["count"] == 3 and r["totals"]["income"] == 1000 and r["totals"]["net"] == 580
    assert [(c["category"], c["total"]) for c in r["categories"]] == [("Rent", 300), ("Fun", 70), ("Food", 50)]
    assert r["previous"]["label"] == "2025" and r["previous"]["spent"] == 100 and r["previous"]["change_pct"] == 320.0
    assert [p["total"] for p in r["periods"]][:3] == [0, 350, 70] and len(r["periods"]) == 12
    assert r["peak"]["label"] == "February"
    assert [s["shop"] for s in r["top_shops"]] == ["Landlord", "SF", "ICA"]
    assert r["largest"][0]["description"] == "Rent" and r["category_matrix"]["rows"][0]["category"] == "Rent"
    assert r["currency"] and r["unconverted_count"] == 0 and 2026 in r["years"] and 2025 in r["years"]


def test_a_month_report(client, make_user):
    _, h = make_user()
    seed(client, h)
    r = get(client, h, year=2026, month=2).json()
    assert (r["kind"], r["label"]) == ("month", "February 2026") and r["totals"]["spent"] == 350 and r["totals"]["count"] == 2
    assert len(r["periods"]) == 28 and r["category_matrix"] is None
    assert r["previous"]["label"] == "January 2026" and r["previous"]["spent"] == 0 and r["previous"]["change_pct"] is None
    assert r["peak"]["label"] == "11 February"


def test_the_year_defaults_to_this_year(client, make_user):
    _, h = make_user()
    assert get(client, h).json()["year"] == THIS_YEAR
    assert get(client, h).json()["years"] == [THIS_YEAR]
    add(client, h, 5, "2023-05-05")
    assert get(client, h).json()["years"] == [THIS_YEAR, 2023]


def test_an_empty_report(client, make_user):
    _, h = make_user()
    r = get(client, h, year=2026).json()
    assert r["totals"]["spent"] == 0 and r["categories"] == [] and r["peak"] is None and r["top_shops"] == [] and r["largest"] == []


def test_shop_merges_apply(client, make_user):
    _, h = make_user()
    add(client, h, 10, "2026-02-01", shop="ICA Maxi"); add(client, h, 20, "2026-02-02", shop="ica")
    m = client.post("/api/aliases/merge", json={"kind": "shop", "names": ["ICA Maxi", "ica"], "canonical": "ICA"}, headers=h)
    assert m.status_code in (200, 201), m.text
    shops = get(client, h, year=2026).json()["top_shops"]
    assert [(s["shop"], s["total"], s["count"]) for s in shops] == [("ICA", 30, 2)]


def test_year_and_month_are_checked(client, make_user):
    _, h = make_user()
    for params in ({"year": 1999}, {"year": 2101}, {"year": "x"}, {"month": 0}, {"month": 13}, {"year": 2026, "month": "x"}):
        assert client.get(URL, params=params, headers=h).status_code == 422, params
        assert client.get(URL + "/export", params=params, headers=h).status_code == 422, params


def test_scope_personal_leaves_space_expenses_out(client, make_user):
    _, h = make_user()
    gid = client.post("/api/groups", json={"name": "G", "space_type": "household"}, headers=h).json()["id"]
    add(client, h, 50, "2026-02-01")
    add(client, h, 70, "2026-02-02", group_id=gid)
    assert get(client, h, year=2026, scope="personal").json()["totals"]["spent"] == 50
    assert get(client, h, year=2026, scope="all").json()["totals"]["spent"] == 120


def test_needs_a_login(client):
    assert client.get(URL).status_code in (401, 403)
    assert client.get(URL + "/export").status_code in (401, 403)


def test_each_person_sees_only_their_own_spending(client, make_user):
    _, a = make_user(); _, b = make_user()
    seed(client, a)
    r = get(client, b, year=2026).json()
    assert r["totals"]["spent"] == 0 and r["totals"]["income"] == 0 and r["top_shops"] == [] and r["years"] == [THIS_YEAR]
    wb = openpyxl.load_workbook(io.BytesIO(get(client, b, URL + "/export", year=2026, format="xlsx").content))
    assert [r for r in wb["Categories"].iter_rows(min_row=2, values_only=True)] == []
    assert [r for r in wb["Top shops"].iter_rows(min_row=2, values_only=True)] == []


def test_export_excel(client, make_user):
    _, h = make_user()
    seed(client, h)
    r = get(client, h, URL + "/export", year=2026, format="xlsx")
    assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert r.headers["content-disposition"] == "attachment; filename=report-2026.xlsx" and r.headers["cache-control"] == "no-store"
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[:3] == ["Summary", "Categories", "By month"]
    values = {row[0]: row[1] for row in wb["Summary"].iter_rows(values_only=True) if row[0]}
    assert values["Total spent"] == 420 and values["Income"] == 1000


def test_export_pdf(client, make_user):
    _, h = make_user()
    seed(client, h)
    r = get(client, h, URL + "/export", year=2026, month=2, format="pdf")
    assert r.headers["content-type"] == "application/pdf" and r.headers["content-disposition"] == "attachment; filename=report-2026-02.pdf"
    assert r.content.startswith(b"%PDF-")


def test_export_defaults_to_excel_for_this_year_and_rejects_other_formats(client, make_user):
    _, h = make_user()
    r = get(client, h, URL + "/export")
    assert r.headers["content-disposition"] == f"attachment; filename=report-{THIS_YEAR}.xlsx"
    for bad in ("csv", "docx", "XLSX", ""):
        assert client.get(URL + "/export", params={"format": bad}, headers=h).status_code == 422, bad


def test_the_export_matches_the_screen(client, make_user):
    _, h = make_user()
    seed(client, h)
    screen = get(client, h, year=2026).json()
    wb = openpyxl.load_workbook(io.BytesIO(get(client, h, URL + "/export", year=2026, format="xlsx").content))
    cats = [(r[0], r[1]) for r in wb["Categories"].iter_rows(min_row=2, values_only=True) if r[0] and r[0] != "Total"]
    assert cats == [(c["category"], c["total"]) for c in screen["categories"]]
