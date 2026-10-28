"""Excel and PDF files for a spending report (the dict built by `app.logic.report.build_report`).

The files only lay out what the report already holds, so they always agree with the screen. Text that
came from the user is neutralised in the spreadsheet (so "=HYPERLINK(...)" cannot run) and made safe
for the PDF's built-in font (characters it cannot draw become "?").
"""
from __future__ import annotations

import io
from datetime import date
from decimal import Decimal
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.graphics.charts.barcharts import HorizontalBarChart, VerticalBarChart
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.parsing import escape_formula

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_MEDIA_TYPE = "application/pdf"

ACCENT = "6D7BFF"
MONEY_FORMAT = "#,##0.00"
CHART_CATEGORIES = 15      # the chart shows this many categories, the rest as one "Other" bar
NAME_LIMIT = 28


def file_stem(report: dict) -> str:
    return f"report-{report['year']}-{report['month']:02d}" if report.get("month") else f"report-{report['year']}"


def _num(value) -> float:
    return float(value) if isinstance(value, (Decimal, int, float)) else 0.0


def _text(value):
    return escape_formula(value) if isinstance(value, str) else value


# ── Excel ────────────────────────────────────────────────────────────────────

def _header(ws, row: int, labels: list[str]) -> None:
    for col, label in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=ACCENT)
        cell.alignment = Alignment(horizontal="left" if col == 1 else "right")


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _bar_chart(ws, title: str, labels: Reference, values: Reference, anchor: str, horizontal: bool, height: float, width: float = 18) -> None:
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.title = title
    chart.legend = None
    chart.height, chart.width = height, width
    chart.add_data(values, titles_from_data=True)
    chart.set_categories(labels)
    chart.series[0].graphicalProperties.solidFill = ACCENT
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    chart.dataLabels.showSerName = chart.dataLabels.showCatName = chart.dataLabels.showLegendKey = False
    if horizontal:
        chart.x_axis.scaling.orientation = "maxMin"     # biggest category at the top
    chart.x_axis.delete = chart.y_axis.delete = False
    ws.add_chart(chart, anchor)


def report_xlsx(report: dict, currency: str) -> bytes:
    totals, previous = report["totals"], report["previous"]
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"Spending report {report['label']}"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"All amounts in {currency}"
    lines: list[tuple[str, object, str | None]] = [
        ("Total spent", _num(totals["spent"]), MONEY_FORMAT),
        ("Number of expenses", totals["count"], "0"),
        ("Income", _num(totals["income"]), MONEY_FORMAT),
        ("Net (income minus spending)", _num(totals["net"]), MONEY_FORMAT),
        ("Savings rate", None if totals["savings_rate_pct"] is None else totals["savings_rate_pct"] / 100, "0.0%"),
        ("Average per day", _num(totals["avg_per_day"]), MONEY_FORMAT),
    ]
    if totals["avg_per_month"] is not None:
        lines.append(("Average per month", _num(totals["avg_per_month"]), MONEY_FORMAT))
    lines += [
        (f"Spent in {previous['label']}", _num(previous["spent"]), MONEY_FORMAT),
        (f"Change from {previous['label']}", _num(previous["change"]), MONEY_FORMAT),
        ("Change in percent", None if previous["change_pct"] is None else previous["change_pct"] / 100, "0.0%"),
    ]
    if report["peak"]:
        lines.append((f"Highest {'day' if report['kind'] == 'month' else 'month'}: {report['peak']['label']}", _num(report["peak"]["total"]), MONEY_FORMAT))
    for i, (label, value, fmt) in enumerate(lines, start=4):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        cell = ws.cell(row=i, column=2, value="–" if value is None else value)
        if value is not None and fmt:
            cell.number_format = fmt
        cell.alignment = Alignment(horizontal="right")
    _widths(ws, [38, 18])

    ws = wb.create_sheet("Categories")
    _header(ws, 1, ["Category", f"Amount ({currency})", "Share", "Expenses"])
    for r, c in enumerate(report["categories"], start=2):
        ws.cell(row=r, column=1, value=_text(c["category"]))
        ws.cell(row=r, column=2, value=_num(c["total"])).number_format = MONEY_FORMAT
        ws.cell(row=r, column=3, value=c["share_pct"] / 100).number_format = "0.0%"
        ws.cell(row=r, column=4, value=c["count"])
    n = len(report["categories"])
    if n:
        total_row = n + 2
        ws.cell(row=total_row, column=1, value="Total").font = Font(bold=True)
        ws.cell(row=total_row, column=2, value=_num(totals["spent"])).number_format = MONEY_FORMAT
        ws.cell(row=total_row, column=2).font = Font(bold=True)
        ws.cell(row=total_row, column=4, value=totals["count"]).font = Font(bold=True)
        _bar_chart(ws, "Spending by category", Reference(ws, min_col=1, min_row=2, max_row=n + 1), Reference(ws, min_col=2, min_row=1, max_row=n + 1),
                   "F2", horizontal=True, height=max(8.0, 1.0 * n + 4))
    _widths(ws, [30, 18, 10, 12])
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("By day" if report["kind"] == "month" else "By month")
    _header(ws, 1, ["Day" if report["kind"] == "month" else "Month", f"Amount ({currency})", "Expenses"])
    for r, p in enumerate(report["periods"], start=2):
        ws.cell(row=r, column=1, value=p["key"] if report["kind"] == "month" else p["label"])
        ws.cell(row=r, column=2, value=_num(p["total"])).number_format = MONEY_FORMAT
        ws.cell(row=r, column=3, value=p["count"])
    count = len(report["periods"])
    _bar_chart(ws, "Spending by day" if report["kind"] == "month" else "Spending by month", Reference(ws, min_col=1, min_row=2, max_row=count + 1),
               Reference(ws, min_col=2, min_row=1, max_row=count + 1), "E2", horizontal=False, height=9, width=22)
    _widths(ws, [14, 18, 12])
    ws.freeze_panes = "A2"

    matrix = report.get("category_matrix")
    if matrix:
        ws = wb.create_sheet("Category by month")
        _header(ws, 1, ["Category"] + [p["label"] for p in report["periods"]] + ["Total"])
        for r, row in enumerate(matrix["rows"], start=2):
            ws.cell(row=r, column=1, value=_text(row["category"]))
            for col, v in enumerate(row["values"], start=2):
                ws.cell(row=r, column=col, value=_num(v)).number_format = MONEY_FORMAT
            ws.cell(row=r, column=14, value=_num(row["total"])).number_format = MONEY_FORMAT
            ws.cell(row=r, column=14).font = Font(bold=True)
        last = len(matrix["rows"]) + 2
        ws.cell(row=last, column=1, value="Total").font = Font(bold=True)
        for col in range(2, 14):
            cell = ws.cell(row=last, column=col, value=_num(report["periods"][col - 2]["total"]))
            cell.number_format, cell.font = MONEY_FORMAT, Font(bold=True)
        ws.cell(row=last, column=14, value=_num(totals["spent"])).number_format = MONEY_FORMAT
        ws.cell(row=last, column=14).font = Font(bold=True)
        _widths(ws, [24] + [11] * 12 + [14])
        ws.freeze_panes = "B2"

    ws = wb.create_sheet("Top shops")
    _header(ws, 1, ["Shop", f"Amount ({currency})", "Share", "Expenses"])
    for r, s in enumerate(report["top_shops"], start=2):
        ws.cell(row=r, column=1, value=_text(s["shop"]))
        ws.cell(row=r, column=2, value=_num(s["total"])).number_format = MONEY_FORMAT
        ws.cell(row=r, column=3, value=s["share_pct"] / 100).number_format = "0.0%"
        ws.cell(row=r, column=4, value=s["count"])
    _widths(ws, [30, 18, 10, 12])

    ws = wb.create_sheet("Largest expenses")
    _header(ws, 1, ["Date", "Category", "Item", "Shop", f"Amount ({currency})"])
    for r, x in enumerate(report["largest"], start=2):
        ws.cell(row=r, column=1, value=date.fromisoformat(x["date"])).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=2, value=_text(x["category"]))
        ws.cell(row=r, column=3, value=_text(x["description"]))
        ws.cell(row=r, column=4, value=_text(x["shop"]))
        ws.cell(row=r, column=5, value=_num(x["amount"])).number_format = MONEY_FORMAT
    _widths(ws, [13, 20, 30, 22, 18])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ── PDF ──────────────────────────────────────────────────────────────────────

def _pdf_safe(value) -> str:
    """Text the built-in Helvetica can draw: anything outside Latin-1/Windows-1252 becomes "?"."""
    return str(value or "").encode("cp1252", "replace").decode("cp1252")


def _p(value, style) -> Paragraph:
    return Paragraph(escape(_pdf_safe(value)), style)


def _money(value) -> str:
    return f"{_num(value):,.2f}"


def _short(value: str, limit: int = NAME_LIMIT) -> str:
    value = _pdf_safe(value)
    return value if len(value) <= limit else value[: limit - 1] + "…".encode("cp1252", "replace").decode("cp1252")


def _category_chart(categories: list[dict], width: float) -> Drawing:
    shown = categories[:CHART_CATEGORIES]
    rest = categories[CHART_CATEGORIES:]
    names = [_short(c["category"]) for c in shown]
    values = [_num(c["total"]) for c in shown]
    if rest:
        names.append(f"Other ({len(rest)})")
        values.append(sum(_num(c["total"]) for c in rest))
    bar_h = 15
    height = len(values) * bar_h + 20
    d = Drawing(width, height)
    chart = HorizontalBarChart()
    chart.x, chart.y, chart.width, chart.height = 120, 10, width - 130, height - 16
    chart.data = [list(reversed(values))]            # the first category ends up at the top
    chart.categoryAxis.categoryNames = list(reversed(names))
    chart.categoryAxis.labels.fontName = "Helvetica"
    chart.categoryAxis.labels.fontSize = 8
    chart.categoryAxis.labels.boxAnchor = "e"
    chart.categoryAxis.visibleTicks = 0
    top = max(max(values), 0.0)
    chart.valueAxis.valueMin = min(0.0, min(values))
    chart.valueAxis.valueMax = top * 1.22 if top else 1
    chart.valueAxis.visible = False
    chart.valueAxis.visibleGrid = False
    chart.bars[0].fillColor = colors.HexColor("#" + ACCENT)
    chart.bars[0].strokeColor = None
    chart.barWidth = bar_h * 0.6
    chart.groupSpacing = bar_h * 0.4
    chart.barLabelFormat = lambda v: f"{v:,.2f}"
    chart.barLabels.fontName = "Helvetica"
    chart.barLabels.fontSize = 7.5
    chart.barLabels.nudge = 4
    chart.barLabels.boxAnchor = "w"
    d.add(chart)
    return d


def _period_chart(periods: list[dict], width: float, daily: bool) -> Drawing:
    values = [_num(p["total"]) for p in periods]
    d = Drawing(width, 150)
    chart = VerticalBarChart()
    chart.x, chart.y, chart.width, chart.height = 40, 22, width - 55, 110
    chart.data = [values]
    chart.categoryAxis.categoryNames = [p["label"] for p in periods]
    chart.categoryAxis.labels.fontName = "Helvetica"
    chart.categoryAxis.labels.fontSize = 6.5 if daily else 8
    chart.categoryAxis.visibleTicks = 0
    top = max(max(values), 0.0)
    chart.valueAxis.valueMin = min(0.0, min(values))
    chart.valueAxis.valueMax = top * 1.15 if top else 1
    chart.valueAxis.labels.fontName = "Helvetica"
    chart.valueAxis.labels.fontSize = 7
    chart.valueAxis.labelTextFormat = lambda v: f"{v:,.0f}"
    chart.valueAxis.gridStrokeColor = colors.HexColor("#E6E8F2")
    chart.valueAxis.visibleGrid = True
    chart.bars[0].fillColor = colors.HexColor("#" + ACCENT)
    chart.bars[0].strokeColor = None
    chart.barSpacing = 0
    chart.groupSpacing = 2 if daily else 6
    if not daily:
        chart.barLabelFormat = lambda v: f"{v:,.0f}" if v else ""
        chart.barLabels.fontName = "Helvetica"
        chart.barLabels.fontSize = 6.5
        chart.barLabels.nudge = 6
    d.add(chart)
    return d


def _table(rows: list[list], col_widths: list[float], right_cols: tuple[int, ...] = ()) -> Table:
    t = Table(rows, colWidths=col_widths, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + ACCENT)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F6FB")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    style += [("ALIGN", (c, 0), (c, -1), "RIGHT") for c in right_cols]
    t.setStyle(TableStyle(style))
    return t


def report_pdf(report: dict, currency: str, generated: date | None = None) -> bytes:
    generated = generated or date.today()
    totals, previous = report["totals"], report["previous"]
    styles = getSampleStyleSheet()
    title = ParagraphStyle("t", parent=styles["Title"], alignment=TA_LEFT, fontSize=20, spaceAfter=2)
    meta = ParagraphStyle("m", parent=styles["Normal"], fontSize=8.5, textColor=colors.HexColor("#5C6577"), spaceAfter=10)
    h2 = ParagraphStyle("h", parent=styles["Heading2"], fontSize=12.5, spaceBefore=12, spaceAfter=6)
    cell = ParagraphStyle("c", parent=styles["Normal"], fontSize=8.5, leading=10.5)
    note = ParagraphStyle("n", parent=styles["Normal"], fontSize=8.5, textColor=colors.HexColor("#5C6577"))

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
        title=_pdf_safe(f"Spending report {report['label']}"), author="Expense Manager",
    )
    width = A4[0] - 36 * mm
    story: list = [
        _p(f"Spending report {report['label']}", title),
        _p(f"All amounts in {currency} · created {generated.day} {generated.strftime('%b %Y')}", meta),
    ]

    change = ""
    if previous["change_pct"] is not None:
        change = f" ({'+' if previous['change'] >= 0 else '-'}{abs(previous['change_pct']):.1f}%)"
    summary = [
        ["Total spent", _money(totals["spent"])],
        ["Number of expenses", str(totals["count"])],
    ]
    if _num(totals["income"]):
        summary += [["Income", _money(totals["income"])], ["Net (income minus spending)", _money(totals["net"])]]
        if totals["savings_rate_pct"] is not None:
            summary.append(["Savings rate", f"{totals['savings_rate_pct']:.1f}%"])
    summary.append(["Average per day", _money(totals["avg_per_day"])])
    if totals["avg_per_month"] is not None:
        summary.append(["Average per month", _money(totals["avg_per_month"])])
    summary.append([f"Spent in {previous['label']}", f"{_money(previous['spent'])}{change}"])
    if report["peak"]:
        summary.append([f"Highest {'day' if report['kind'] == 'month' else 'month'}", f"{_pdf_safe(report['peak']['label'])}: {_money(report['peak']['total'])}"])
    t = Table(summary, colWidths=[width * 0.5, width * 0.5])
    t.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9.5), ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"), ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E6E8F2")), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(t)

    if not report["categories"]:
        story += [Spacer(1, 14), _p("No spending was recorded in this period.", note)]
    else:
        story.append(_p("Spending by category", h2))
        story.append(_category_chart(report["categories"], width))
        rows = [["Category", f"Amount ({currency})", "Share", "Expenses"]]
        rows += [[_p(c["category"], cell), _money(c["total"]), f"{c['share_pct']:.1f}%", str(c["count"])] for c in report["categories"]]
        rows.append(["Total", _money(totals["spent"]), "", str(totals["count"])])
        ct = _table(rows, [width * 0.46, width * 0.22, width * 0.14, width * 0.18], right_cols=(1, 2, 3))
        ct.setStyle(TableStyle([("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"), ("LINEABOVE", (0, -1), (-1, -1), 0.5, colors.HexColor("#99A1B3"))]))
        story += [Spacer(1, 6), ct]

        daily = report["kind"] == "month"
        story.append(_p("Spending by day" if daily else "Spending by month", h2))
        story.append(_period_chart(report["periods"], width, daily))

        if report["top_shops"]:
            story.append(_p("Top shops", h2))
            rows = [["Shop", f"Amount ({currency})", "Share", "Expenses"]]
            rows += [[_p(s["shop"], cell), _money(s["total"]), f"{s['share_pct']:.1f}%", str(s["count"])] for s in report["top_shops"]]
            story.append(_table(rows, [width * 0.46, width * 0.22, width * 0.14, width * 0.18], right_cols=(1, 2, 3)))

        story.append(_p("Largest expenses", h2))
        rows = [["Date", "Item", "Category", f"Amount ({currency})"]]
        for x in report["largest"]:
            what = x["description"] or x["category"]
            if x["shop"]:
                what = f"{what} · {x['shop']}"
            rows.append([x["date"], _p(what, cell), _p(x["category"], cell), _money(x["amount"])])
        story.append(_table(rows, [width * 0.15, width * 0.43, width * 0.22, width * 0.20], right_cols=(3,)))

    doc.build(story)
    return buffer.getvalue()
